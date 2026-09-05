import os
import re
import io
import uuid
import hmac
import hashlib
import time
from pathlib import Path
from datetime import datetime, timezone, timedelta
from typing import Optional, List

from dotenv import load_dotenv
load_dotenv(Path(__file__).parent / ".env")

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from motor.motor_asyncio import AsyncIOMotorClient

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "bonstok")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "changeme123")
NURSE_PASSWORD = os.environ.get("NURSE_PASSWORD", "changeme456")
TOKEN_SECRET = os.environ.get("TOKEN_SECRET", "change-this-secret")

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

app = FastAPI(title="BonStok API")
api_router = APIRouter(prefix="/api")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def new_id():
    return str(uuid.uuid4())


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def fmt_num(n):
    try:
        return f"{float(n):,.0f}".replace(",", ".")
    except Exception:
        return str(n)


def parse_date_range(start: Optional[str], end: Optional[str]):
    start_dt = None
    end_dt = None
    if start:
        try:
            start_dt = datetime.fromisoformat(start).replace(tzinfo=timezone.utc)
        except Exception:
            start_dt = None
    if end:
        try:
            end_dt = datetime.fromisoformat(end).replace(tzinfo=timezone.utc) + timedelta(days=1)
        except Exception:
            end_dt = None
    return start_dt, end_dt


# ---------------- Auth tokens (HMAC, role-aware, no external deps) ----------------

def make_token(role: str) -> str:
    ts = str(int(time.time()))
    sig = hmac.new(TOKEN_SECRET.encode(), f"{role}.{ts}".encode(), hashlib.sha256).hexdigest()
    return f"{role}.{ts}.{sig}"


def verify_token(token: str) -> Optional[str]:
    """Returns the role encoded in the token if valid & not expired, else None."""
    try:
        role, ts, sig = token.split(".", 2)
        expected = hmac.new(TOKEN_SECRET.encode(), f"{role}.{ts}".encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(sig, expected):
            return None
        if time.time() - int(ts) > 7 * 24 * 3600:
            return None
        return role
    except Exception:
        return None


async def require_admin(authorization: Optional[str] = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Tidak ada token")
    role = verify_token(authorization.split(" ", 1)[1])
    if role != "admin":
        raise HTTPException(status_code=401, detail="Token tidak valid atau kedaluwarsa")
    return True


async def require_nurse(authorization: Optional[str] = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Tidak ada token")
    role = verify_token(authorization.split(" ", 1)[1])
    if role not in ("admin", "nurse"):
        raise HTTPException(status_code=401, detail="Token tidak valid atau kedaluwarsa")
    return role


# ---------------- Models ----------------

class LoginPayload(BaseModel):
    password: str


class ItemPayload(BaseModel):
    name: str
    barcode: str
    unit: str = "pcs"
    category: Optional[str] = None
    current_stock: float = 0
    min_stock: float = 0
    photo: Optional[str] = None


class BonItemLine(BaseModel):
    item_id: str
    qty: float


class BonRequestPayload(BaseModel):
    requester_name: str
    room_id: Optional[str] = None
    room: str
    items: List[BonItemLine]
    note: Optional[str] = None


class MedicinePayload(BaseModel):
    name: str
    unit: str = "pcs"
    category: Optional[str] = None
    current_stock: float = 0
    min_stock: float = 0


class MedicineTransactionPayload(BaseModel):
    medicine_id: str
    type: str  # "masuk" | "keluar"
    qty: float
    nurse_name: str
    note: Optional[str] = None


class RoomPayload(BaseModel):
    name: str


class StockInPayload(BaseModel):
    item_id: Optional[str] = None
    name: Optional[str] = None
    unit: str = "pcs"
    qty: float
    photo: Optional[str] = None
    note: Optional[str] = None


# ---------------- Public: items (gudang) ----------------

@api_router.get("/health")
async def health():
    return {"ok": True, "time": now_iso()}


@api_router.get("/items")
async def list_items(search: Optional[str] = None):
    q = {}
    if search:
        q = {"$or": [
            {"name": {"$regex": search, "$options": "i"}},
            {"barcode": {"$regex": search, "$options": "i"}},
        ]}
    items = await db.items.find(q, {"_id": 0}).sort("name", 1).to_list(500)
    return items


@api_router.get("/items/by-barcode/{barcode}")
async def get_item_by_barcode(barcode: str):
    item = await db.items.find_one({"barcode": barcode}, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Barang dengan barcode ini tidak ditemukan")
    return item


# ---------------- Public: rooms ----------------

@api_router.get("/rooms")
async def list_rooms(search: Optional[str] = None):
    q = {}
    if search:
        q = {"name": {"$regex": search, "$options": "i"}}
    rooms = await db.rooms.find(q, {"_id": 0}).sort("name", 1).to_list(500)
    return rooms


@api_router.get("/rooms/by-barcode/{barcode}")
async def get_room_by_barcode(barcode: str):
    room = await db.rooms.find_one({"barcode": barcode}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Ruangan dengan barcode ini tidak ditemukan")
    return room


# ---------------- Public: bon request ----------------

@api_router.post("/bon")
async def submit_bon(payload: BonRequestPayload):
    if not payload.items:
        raise HTTPException(status_code=400, detail="Keranjang bon kosong")

    room_name = payload.room
    room_id = payload.room_id
    if payload.room_id:
        room = await db.rooms.find_one({"id": payload.room_id})
        if not room:
            raise HTTPException(status_code=404, detail="Ruangan tidak ditemukan")
        room_name = room["name"]

    if not room_name or not room_name.strip():
        raise HTTPException(status_code=400, detail="Ruangan wajib diisi")

    lines = []
    for line in payload.items:
        item = await db.items.find_one({"id": line.item_id})
        if not item:
            raise HTTPException(status_code=404, detail=f"Barang tidak ditemukan: {line.item_id}")
        if line.qty <= 0:
            raise HTTPException(status_code=400, detail="Jumlah harus lebih dari 0")
        lines.append({
            "item_id": item["id"],
            "item_name": item["name"],
            "barcode": item["barcode"],
            "unit": item.get("unit", "pcs"),
            "qty": line.qty,
        })

    doc = {
        "id": new_id(),
        "requester_name": payload.requester_name,
        "room_id": room_id,
        "room": room_name,
        "items": lines,
        "note": payload.note,
        "status": "pending",
        "requested_at": now_iso(),
        "decided_at": None,
        "decided_note": None,
    }
    await db.bon_requests.insert_one(doc)
    doc.pop("_id", None)
    return doc


# ---------------- Nota Dinas (Word export) ----------------

def _remove_table_borders(table):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def build_nota_dinas_docx(bon):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)

    def para(text="", size=12, bold=False, align="center", underline=False, space_after=2):
        p = doc.add_paragraph()
        p.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }[align]
        if text:
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(size)
            r.bold = bold
            r.underline = underline
        p.paragraph_format.space_after = Pt(space_after)
        return p

    para("KEMENTERIAN IMIGRASI DAN PEMASYARAKATAN REPUBLIK INDONESIA", 13)
    para("DIREKTORAT JENDERAL PEMASYARAKATAN", 13)
    para("KANTOR WILAYAH KALIMANTAN TENGAH", 13)
    para("LEMBAGA PEMASYARAKATAN KELAS IIA PALANGKA RAYA", 13, bold=True)
    para("Jalan Tjilik Riwut km. 40, Palangka Raya 73221", 11)
    para("Laman: lapaspalangkaraya.kemenkumham.go.id, Pos-el: lapaspalangkaraya2a@gmail.com", 11, space_after=12)

    para("NOTA DINAS", 14, bold=True, underline=True, space_after=0)
    para("NOMOR : WP.17.PAS.1.UM.03.04- ", 12, space_after=14)

    dots = "…………………………………………………………"

    t0 = doc.add_table(rows=5, cols=3)
    labels = ["Kepada", "Dari", "Hal", "Lampiran", "Tanggal"]
    for i, lab in enumerate(labels):
        t0.cell(i, 0).text = lab
        t0.cell(i, 1).text = ":"
        t0.cell(i, 2).text = dots
        for c in (0, 1, 2):
            for p in t0.cell(i, c).paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(12)
    _remove_table_borders(t0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    para(
        "\tBersama ini dengan hormat kami mohon untuk memenuhi kebutuhan barang gudang "
        "sebagaimana rincian terlampir, guna menunjang kelancaran operasional pada "
        "Lembaga Pemasyarakatan Kelas IIA Palangka Raya.",
        12, align="justify", space_after=10,
    )
    para("Demikian kami sampaikan, atas perhatiannya diucapkan terima kasih.", 12, align="justify", space_after=24)

    t1 = doc.add_table(rows=1, cols=2)
    t1.cell(0, 0).text = ""
    t1.cell(0, 1).text = "Kepala ..................\n\n\n\n\n\n(.........................................)"
    for p in t1.cell(0, 1).paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _remove_table_borders(t1)

    doc.add_paragraph()
    para("RENCANA KEBUTUHAN", 13, bold=True, space_after=10)

    items = bon.get("items") or []
    n = max(len(items), 1)
    t2 = doc.add_table(rows=n + 1, cols=4)
    t2.style = "Table Grid"
    headers = ["No", "NAMA BARANG", "JUMLAH KEBUTUHAN", "REALISASI KEBUTUHAN"]
    for c, h in enumerate(headers):
        cell = t2.cell(0, c)
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(11)

    if items:
        for i, it in enumerate(items, 1):
            t2.cell(i, 0).text = str(i)
            t2.cell(i, 1).text = it.get("item_name", "")
            t2.cell(i, 2).text = f"{fmt_num(it.get('qty', 0))} {it.get('unit', '')}".strip()
            t2.cell(i, 3).text = "…………………………."
            for c in (0, 1, 2, 3):
                for p in t2.cell(i, c).paragraphs:
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(11)
    else:
        t2.cell(1, 0).text = "1"
        for c in (1, 2, 3):
            t2.cell(1, c).text = "…………………………."

    doc.add_paragraph()
    t3 = doc.add_table(rows=2, cols=1)
    t3.cell(0, 0).text = "Catatan Kepala :"
    t3.cell(1, 0).text = ""
    _remove_table_borders(t3)

    doc.add_paragraph()
    t4 = doc.add_table(rows=2, cols=2)
    t4.cell(0, 0).text = (
        "Kepala Sub Bagian Tata Usaha,\n\n\n\n\n\n\nHarjono\nNIP. 196907101991031001"
    )
    t4.cell(0, 1).text = (
        "Palangka Raya, ……………..\n\nKepala Urusan Umum,\n\n\n\n\n\nBoby Rodiarta\nNIP. 198402272010121001"
    )
    merged = t4.cell(1, 0).merge(t4.cell(1, 1))
    merged.text = "Mengetahui,\nKalapas,\n\n\n\n\nHisam Wibowo\nNIP. 197411111997031001"
    for row in t4.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
    _remove_table_borders(t4)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@api_router.get("/bon/{bon_id}/nota-dinas")
async def bon_nota_dinas(bon_id: str):
    bon = await db.bon_requests.find_one({"id": bon_id}, {"_id": 0})
    if not bon:
        raise HTTPException(status_code=404, detail="Bon tidak ditemukan")
    buf = build_nota_dinas_docx(bon)
    filename = f"nota-dinas-bon-{bon_id[:8]}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------- Public: medicines (klinik) — perawat login required ----------------

@api_router.get("/medicines")
async def list_medicines(search: Optional[str] = None, _: str = Depends(require_nurse)):
    q = {}
    if search:
        q = {"name": {"$regex": search, "$options": "i"}}
    meds = await db.medicines.find(q, {"_id": 0}).sort("name", 1).to_list(500)
    return meds


@api_router.post("/medicine-transactions")
async def create_medicine_transaction(payload: MedicineTransactionPayload, _: str = Depends(require_nurse)):
    if payload.type not in ("masuk", "keluar"):
        raise HTTPException(status_code=400, detail="Tipe transaksi tidak valid")
    if payload.qty <= 0:
        raise HTTPException(status_code=400, detail="Jumlah harus lebih dari 0")

    med = await db.medicines.find_one({"id": payload.medicine_id})
    if not med:
        raise HTTPException(status_code=404, detail="Obat tidak ditemukan")

    delta = payload.qty if payload.type == "masuk" else -payload.qty
    new_stock = med.get("current_stock", 0) + delta
    if new_stock < 0:
        raise HTTPException(status_code=400, detail="Stok obat tidak cukup")

    await db.medicines.update_one({"id": payload.medicine_id}, {"$set": {"current_stock": new_stock}})

    doc = {
        "id": new_id(),
        "medicine_id": payload.medicine_id,
        "medicine_name": med["name"],
        "type": payload.type,
        "qty": payload.qty,
        "nurse_name": payload.nurse_name,
        "note": payload.note,
        "created_at": now_iso(),
    }
    await db.medicine_transactions.insert_one(doc)
    doc.pop("_id", None)
    return doc


# ---------------- Auth: admin & perawat login ----------------

@api_router.post("/admin/login")
async def admin_login(payload: LoginPayload):
    if not hmac.compare_digest(payload.password, ADMIN_PASSWORD):
        raise HTTPException(status_code=401, detail="Password salah")
    return {"token": make_token("admin")}


@api_router.post("/nurse/login")
async def nurse_login(payload: LoginPayload):
    if not hmac.compare_digest(payload.password, NURSE_PASSWORD):
        raise HTTPException(status_code=401, detail="Password salah")
    return {"token": make_token("nurse")}


# ---------------- Admin: items ----------------

@api_router.get("/admin/items")
async def admin_list_items(_: bool = Depends(require_admin)):
    items = await db.items.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return items


@api_router.post("/admin/items")
async def admin_create_item(payload: ItemPayload, _: bool = Depends(require_admin)):
    existing = await db.items.find_one({"barcode": payload.barcode})
    if existing:
        raise HTTPException(status_code=400, detail="Barcode sudah terdaftar")
    doc = payload.model_dump()
    doc["id"] = new_id()
    doc["created_at"] = now_iso()
    await db.items.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/admin/items/{item_id}")
async def admin_update_item(item_id: str, payload: ItemPayload, _: bool = Depends(require_admin)):
    existing = await db.items.find_one({"id": item_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
    dup = await db.items.find_one({"barcode": payload.barcode, "id": {"$ne": item_id}})
    if dup:
        raise HTTPException(status_code=400, detail="Barcode sudah dipakai barang lain")
    doc = payload.model_dump()
    doc["updated_at"] = now_iso()
    await db.items.update_one({"id": item_id}, {"$set": doc})
    updated = await db.items.find_one({"id": item_id}, {"_id": 0})
    return updated


@api_router.delete("/admin/items/{item_id}")
async def admin_delete_item(item_id: str, _: bool = Depends(require_admin)):
    await db.items.delete_one({"id": item_id})
    return {"ok": True}


# ---------------- Admin: stock in & item transactions (persediaan) ----------------

@api_router.post("/admin/items/stock-in")
async def admin_stock_in(payload: StockInPayload, _: bool = Depends(require_admin)):
    if payload.qty <= 0:
        raise HTTPException(status_code=400, detail="Jumlah harus lebih dari 0")

    item = None
    if payload.item_id:
        item = await db.items.find_one({"id": payload.item_id})
        if not item:
            raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
    elif payload.name and payload.name.strip():
        item = await db.items.find_one({"name": {"$regex": f"^{re.escape(payload.name.strip())}$", "$options": "i"}})
        if not item:
            barcode = f"AUTO-{new_id()[:8].upper()}"
            while await db.items.find_one({"barcode": barcode}):
                barcode = f"AUTO-{new_id()[:8].upper()}"
            doc = {
                "id": new_id(),
                "name": payload.name.strip(),
                "barcode": barcode,
                "unit": payload.unit or "pcs",
                "category": None,
                "current_stock": 0,
                "min_stock": 0,
                "photo": payload.photo,
                "created_at": now_iso(),
            }
            await db.items.insert_one(doc)
            item = doc
    else:
        raise HTTPException(status_code=400, detail="Pilih barang yang sudah ada atau isi nama barang baru")

    await db.items.update_one({"id": item["id"]}, {"$inc": {"current_stock": payload.qty}})

    tx = {
        "id": new_id(),
        "item_id": item["id"],
        "item_name": item["name"],
        "type": "masuk",
        "qty": payload.qty,
        "unit": payload.unit or item.get("unit", "pcs"),
        "photo": payload.photo,
        "note": payload.note,
        "created_at": now_iso(),
    }
    await db.item_transactions.insert_one(tx)
    tx.pop("_id", None)
    return tx


@api_router.get("/admin/item-transactions")
async def admin_list_item_transactions(item_id: Optional[str] = None, _: bool = Depends(require_admin)):
    q = {}
    if item_id:
        q["item_id"] = item_id
    items = await db.item_transactions.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return items


# ---------------- Admin: rooms ----------------

@api_router.get("/admin/rooms")
async def admin_list_rooms(_: bool = Depends(require_admin)):
    rooms = await db.rooms.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return rooms


@api_router.post("/admin/rooms")
async def admin_create_room(payload: RoomPayload, _: bool = Depends(require_admin)):
    barcode = f"ROOM-{new_id()[:8].upper()}"
    while await db.rooms.find_one({"barcode": barcode}):
        barcode = f"ROOM-{new_id()[:8].upper()}"
    doc = {
        "id": new_id(),
        "name": payload.name,
        "barcode": barcode,
        "created_at": now_iso(),
    }
    await db.rooms.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/admin/rooms/{room_id}")
async def admin_update_room(room_id: str, payload: RoomPayload, _: bool = Depends(require_admin)):
    existing = await db.rooms.find_one({"id": room_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Ruangan tidak ditemukan")
    await db.rooms.update_one({"id": room_id}, {"$set": {"name": payload.name, "updated_at": now_iso()}})
    updated = await db.rooms.find_one({"id": room_id}, {"_id": 0})
    return updated


@api_router.delete("/admin/rooms/{room_id}")
async def admin_delete_room(room_id: str, _: bool = Depends(require_admin)):
    await db.rooms.delete_one({"id": room_id})
    return {"ok": True}


# ---------------- Admin: bon requests ----------------

@api_router.get("/admin/bon")
async def admin_list_bon(status: Optional[str] = None, _: bool = Depends(require_admin)):
    q = {}
    if status:
        q["status"] = status
    items = await db.bon_requests.find(q, {"_id": 0}).sort("requested_at", -1).to_list(500)
    return items


@api_router.post("/admin/bon/{bon_id}/approve")
async def admin_approve_bon(bon_id: str, _: bool = Depends(require_admin)):
    bon = await db.bon_requests.find_one({"id": bon_id})
    if not bon:
        raise HTTPException(status_code=404, detail="Bon tidak ditemukan")
    if bon["status"] != "pending":
        raise HTTPException(status_code=400, detail="Bon sudah diproses")

    for line in bon["items"]:
        item = await db.items.find_one({"id": line["item_id"]})
        if not item:
            raise HTTPException(status_code=404, detail=f"Barang {line['item_name']} tidak ditemukan")
        if item.get("current_stock", 0) < line["qty"]:
            raise HTTPException(status_code=400, detail=f"Stok {line['item_name']} tidak cukup")

    for line in bon["items"]:
        await db.items.update_one(
            {"id": line["item_id"]},
            {"$inc": {"current_stock": -line["qty"]}},
        )
        await db.item_transactions.insert_one({
            "id": new_id(),
            "item_id": line["item_id"],
            "item_name": line["item_name"],
            "type": "keluar",
            "qty": line["qty"],
            "unit": line.get("unit", "pcs"),
            "photo": None,
            "note": f"Bon: {bon['requester_name']} - {bon.get('room', '')}",
            "created_at": now_iso(),
        })

    await db.bon_requests.update_one(
        {"id": bon_id},
        {"$set": {"status": "approved", "decided_at": now_iso()}},
    )
    updated = await db.bon_requests.find_one({"id": bon_id}, {"_id": 0})
    return updated


class DecisionPayload(BaseModel):
    note: Optional[str] = None


@api_router.post("/admin/bon/{bon_id}/reject")
async def admin_reject_bon(bon_id: str, payload: DecisionPayload, _: bool = Depends(require_admin)):
    bon = await db.bon_requests.find_one({"id": bon_id})
    if not bon:
        raise HTTPException(status_code=404, detail="Bon tidak ditemukan")
    if bon["status"] != "pending":
        raise HTTPException(status_code=400, detail="Bon sudah diproses")

    await db.bon_requests.update_one(
        {"id": bon_id},
        {"$set": {"status": "rejected", "decided_at": now_iso(), "decided_note": payload.note}},
    )
    updated = await db.bon_requests.find_one({"id": bon_id}, {"_id": 0})
    return updated


# ---------------- Admin: medicines ----------------

@api_router.get("/admin/medicines")
async def admin_list_medicines(_: bool = Depends(require_admin)):
    meds = await db.medicines.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return meds


@api_router.post("/admin/medicines")
async def admin_create_medicine(payload: MedicinePayload, _: bool = Depends(require_admin)):
    doc = payload.model_dump()
    doc["id"] = new_id()
    doc["created_at"] = now_iso()
    await db.medicines.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/admin/medicines/{medicine_id}")
async def admin_update_medicine(medicine_id: str, payload: MedicinePayload, _: bool = Depends(require_admin)):
    existing = await db.medicines.find_one({"id": medicine_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Obat tidak ditemukan")
    doc = payload.model_dump()
    doc["updated_at"] = now_iso()
    await db.medicines.update_one({"id": medicine_id}, {"$set": doc})
    updated = await db.medicines.find_one({"id": medicine_id}, {"_id": 0})
    return updated


@api_router.delete("/admin/medicines/{medicine_id}")
async def admin_delete_medicine(medicine_id: str, _: bool = Depends(require_admin)):
    await db.medicines.delete_one({"id": medicine_id})
    return {"ok": True}


@api_router.get("/admin/medicine-transactions")
async def admin_list_medicine_transactions(medicine_id: Optional[str] = None, _: bool = Depends(require_admin)):
    q = {}
    if medicine_id:
        q["medicine_id"] = medicine_id
    items = await db.medicine_transactions.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return items


# ---------------- Admin: stats ----------------

@api_router.get("/admin/stats")
async def admin_stats(_: bool = Depends(require_admin)):
    total_items = await db.items.count_documents({})
    low_stock_items = await db.items.count_documents({"$expr": {"$lte": ["$current_stock", "$min_stock"]}})
    total_medicines = await db.medicines.count_documents({})
    low_stock_medicines = await db.medicines.count_documents({"$expr": {"$lte": ["$current_stock", "$min_stock"]}})
    pending_bon = await db.bon_requests.count_documents({"status": "pending"})
    approved_bon = await db.bon_requests.count_documents({"status": "approved"})
    rejected_bon = await db.bon_requests.count_documents({"status": "rejected"})
    total_rooms = await db.rooms.count_documents({})

    return {
        "total_items": total_items,
        "low_stock_items": low_stock_items,
        "total_medicines": total_medicines,
        "low_stock_medicines": low_stock_medicines,
        "pending_bon": pending_bon,
        "approved_bon": approved_bon,
        "rejected_bon": rejected_bon,
        "total_rooms": total_rooms,
    }


# ---------------- Admin: laporan persediaan (export PDF / Excel) ----------------

def build_pdf_report(rows, title_label, period_label):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"<b>Laporan Persediaan — {title_label}</b>", styles["Title"]),
        Paragraph(f"Periode: {period_label}", styles["Normal"]),
        Paragraph(f"Dicetak: {now_iso()[:19].replace('T', ' ')} UTC", styles["Normal"]),
        Spacer(1, 12),
    ]

    data = [["No", "Nama Barang", "Satuan", "Total Masuk", "Total Keluar", "Saldo Akhir"]]
    for i, r in enumerate(rows, 1):
        data.append([str(i), r["name"], r["unit"], fmt_num(r["masuk"]), fmt_num(r["keluar"]), fmt_num(r["saldo"])])

    table = Table(data, repeatRows=1, colWidths=[1.2 * cm, 6.5 * cm, 2 * cm, 3 * cm, 3 * cm, 3 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#171717")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d4d4d4")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    filename = f"laporan-persediaan-{title_label.lower().replace(' ', '-')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                              headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def build_xlsx_report(rows, title_label, period_label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Persediaan"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Laporan Persediaan — {title_label}"
    ws["A1"].font = Font(bold=True, size=14)

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Periode: {period_label}"

    ws.append([])
    headers = ["No", "Nama Barang", "Satuan", "Total Masuk", "Total Keluar", "Saldo Akhir"]
    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, 7):
        c = ws.cell(row=header_row, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="171717")
        c.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows, 1):
        ws.append([i, r["name"], r["unit"], r["masuk"], r["keluar"], r["saldo"]])

    widths = [6, 36, 10, 14, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"laporan-persediaan-{title_label.lower().replace(' ', '-')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/admin/report")
async def admin_report(
    module: str = "items",
    start: Optional[str] = None,
    end: Optional[str] = None,
    format: str = "pdf",
    _: bool = Depends(require_admin),
):
    if module not in ("items", "medicines"):
        raise HTTPException(status_code=400, detail="Modul tidak valid")
    if format not in ("pdf", "xlsx"):
        raise HTTPException(status_code=400, detail="Format tidak valid")

    start_dt, end_dt = parse_date_range(start, end)

    if module == "items":
        master = await db.items.find({}, {"_id": 0}).sort("name", 1).to_list(2000)
        txs = await db.item_transactions.find({}, {"_id": 0}).to_list(20000)
        title_label = "Barang Gudang"
    else:
        master = await db.medicines.find({}, {"_id": 0}).sort("name", 1).to_list(2000)
        txs = await db.medicine_transactions.find({}, {"_id": 0}).to_list(20000)
        title_label = "Obat Klinik"

    totals = {}
    for tx in txs:
        try:
            ts = datetime.fromisoformat(tx["created_at"])
        except Exception:
            continue
        if start_dt and ts < start_dt:
            continue
        if end_dt and ts >= end_dt:
            continue
        key = tx.get("item_id") or tx.get("medicine_id")
        bucket = totals.setdefault(key, {"masuk": 0.0, "keluar": 0.0})
        if tx.get("type") == "masuk":
            bucket["masuk"] += tx.get("qty", 0)
        elif tx.get("type") == "keluar":
            bucket["keluar"] += tx.get("qty", 0)

    rows = []
    for m in master:
        t = totals.get(m["id"], {"masuk": 0.0, "keluar": 0.0})
        rows.append({
            "name": m["name"],
            "unit": m.get("unit", "pcs"),
            "masuk": t["masuk"],
            "keluar": t["keluar"],
            "saldo": m.get("current_stock", 0),
        })

    period_label = f"{start or '-'} s/d {end or '-'}"

    if format == "xlsx":
        return build_xlsx_report(rows, title_label, period_label)
    return build_pdf_report(rows, title_label, period_label)


# ---------------- Startup ----------------

@app.on_event("startup")
async def startup():
    await db.items.create_index("id", unique=True)
    await db.items.create_index("barcode", unique=True)
    await db.medicines.create_index("id", unique=True)
    await db.bon_requests.create_index("id", unique=True)
    await db.bon_requests.create_index("status")
    await db.medicine_transactions.create_index("id", unique=True)
    await db.medicine_transactions.create_index("medicine_id")
    await db.rooms.create_index("id", unique=True)
    await db.rooms.create_index("barcode", unique=True)
    await db.item_transactions.create_index("id", unique=True)
    await db.item_transactions.create_index("item_id")


app.include_router(api_router)
